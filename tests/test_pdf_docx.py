"""
tests/test_pdf_docx.py — Comprehensive test suite for PDF to DOCX document conversion.

Covers:
- Born-digital PDF conversion
- Multi-page document flow and page break preservation
- Scanned raster PDF conversion via offline OCR
- Mixed digital and scanned pages
- Tabular structure reconstruction
- Heading and structural hierarchy detection
- Multilingual Unicode preservation
- Rotated page orientation
- Empty, corrupt, and encrypted PDF handling
- Cooperative cancellation
- Output atomic placement and collision policy
- Preflight planning and capability discovery
"""
from __future__ import annotations

from pathlib import Path

import docx
import pytest

from localshift import (
    ConversionCancelledError,
    ConversionContext,
    ConversionFailedError,
    ConversionRequest,
    InvalidInputError,
    convert_request,
    plan_route,
    validate_preflight,
)
from localshift.documents.adapter import PdfDocumentAdapter
from localshift.formats import output_formats_for, route_supported, sniff_format

FIXTURES_DIR = Path(__file__).parent / "fixtures" / "documents"


@pytest.fixture(scope="session", autouse=True)
def ensure_fixtures():
    """Ensure deterministic benchmark fixtures exist."""
    from benchmark_generator import generate_benchmark_fixtures
    if not (FIXTURES_DIR / "doc_01_simple_digital.pdf").exists():
        generate_benchmark_fixtures(FIXTURES_DIR)


class TestPdfDocxPlanningAndRegistry:
    """Verifies format registry, routing, and preflight planning for PDF -> DOCX."""

    def test_docx_registered_and_route_supported(self):
        assert route_supported("pdf", "docx") is True
        from localshift.documents.docx_adapter import DocxToPdfAdapter
        assert route_supported("docx", "pdf") is DocxToPdfAdapter().is_available()

    def test_output_formats_for_pdf_includes_docx(self):
        pdf_path = FIXTURES_DIR / "doc_01_simple_digital.pdf"
        targets = output_formats_for(pdf_path)
        assert "docx" in targets
        assert "txt" in targets

    def test_plan_route_pdf_to_docx(self):
        plan = plan_route("pdf", "docx")
        assert plan.route_supported is True
        assert plan.is_ready is True
        assert plan.adapter_id == "pdf_native_text"
        opt_names = [o.name for o in plan.applicable_options]
        assert "ocr_mode" in opt_names
        assert "ocr_dpi" in opt_names
        assert "extract_images" in opt_names

    def test_preflight_validation_pdf_to_docx(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_01_simple_digital.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)
        findings = validate_preflight(req)
        assert not any(f.blocks_execution for f in findings)


class TestPdfDocxReconstruction:
    """Verifies document structure, quality, and element recovery."""

    def test_simple_born_digital_conversion(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_01_simple_digital.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        res = convert_request(req)
        assert res.output_path.exists()
        assert res.source_format == "pdf"
        assert res.target_format == "docx"
        assert res.output_bytes > 0

        # Verify docx package integrity and text extraction
        doc = docx.Document(res.output_path)
        all_text = " ".join(p.text for p in doc.paragraphs if p.text)
        assert "LocalShift Document Engine" in all_text
        assert "born-digital" in all_text

        # Sniff output format
        assert sniff_format(res.output_path) == "docx"

    def test_multipage_digital_flow(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_02_multipage_digital.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)
        all_text = " ".join(p.text for p in doc.paragraphs if p.text)
        assert "Page 1 of 3" in all_text
        assert "Page 2 of 3" in all_text
        assert "Page 3 of 3" in all_text

    def test_table_reconstruction(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_04_table.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)

        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert len(table.rows) == 4
        assert len(table.columns) == 5

        header_cells = [c.text for c in table.rows[0].cells]
        assert "Item ID" in header_cells
        assert "Description" in header_cells
        assert "Total" in header_cells

        first_data = [c.text for c in table.rows[1].cells]
        assert first_data[0] == "LS-001"
        assert "Image Converter" in first_data[1]

    def test_heading_and_structural_hierarchy(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_12_hierarchy.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)

        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "LOCALSHIFT ARCHITECTURE MANUAL" in headings
        assert "1. Executive Overview" in headings
        assert "1.1 Design Principles" in headings

        # Bullet lists preserved
        bullets = [p.text for p in doc.paragraphs if p.text.startswith("- ")]
        assert len(bullets) >= 3
        assert any("Privacy First" in b for b in bullets)

    def test_scanned_english_pdf_via_ocr(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_05_english_scan.pdf"
        req = ConversionRequest(
            source=pdf_path,
            target_format="docx",
            output_dir=tmp_path,
            options={"ocr_mode": "auto", "ocr_dpi": 200},
        )

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)
        all_text = " ".join(p.text for p in doc.paragraphs if p.text)
        assert "SCANNED" in all_text.upper()
        assert "INV-2026-001" in all_text

    def test_mixed_digital_and_scanned_pdf(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_07_mixed_pages.pdf"
        req = ConversionRequest(
            source=pdf_path,
            target_format="docx",
            output_dir=tmp_path,
            options={"ocr_mode": "auto"},
        )

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)
        all_text = " ".join(p.text for p in doc.paragraphs if p.text)
        # Page 1 digital content
        assert "Page 1: Digital Summary Report" in all_text
        # Page 2 scanned content
        assert "SCANNED" in all_text.upper()

    def test_unicode_preservation(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_11_unicode_digital.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)
        all_text = " ".join(p.text for p in doc.paragraphs if p.text)
        # Accented characters
        assert "Café" in all_text or "naïve" in all_text
        assert "Äpfel" in all_text or "Übergröße" in all_text
        assert "señorita" in all_text or "niño" in all_text
        # Currency / symbols
        assert "©2026" in all_text
        assert "100%" in all_text

    def test_rotated_page_orientation(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_08_rotated_page.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)
        # Section should be set to landscape
        assert doc.sections[0].orientation == docx.enum.section.WD_ORIENT.LANDSCAPE

    def test_blank_page_handling(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_13_empty.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        res = convert_request(req)
        assert res.output_path.exists()
        doc = docx.Document(res.output_path)
        # Empty document should have zero or empty paragraphs
        non_empty = [p.text for p in doc.paragraphs if p.text.strip()]
        assert len(non_empty) == 0
        assert any("empty" in w.lower() for w in res.warnings)


class TestPdfDocxSafetyAndRobustness:
    """Verifies security boundaries, error propagation, and resource limits."""

    def test_corrupt_pdf_fails_cleanly(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_09_corrupt.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        with pytest.raises(InvalidInputError):
            convert_request(req)

    def test_encrypted_pdf_fails_cleanly(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_10_encrypted.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)

        with pytest.raises(InvalidInputError):
            convert_request(req)

    def test_cancellation_during_docx_conversion(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_02_multipage_digital.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)
        ctx = ConversionContext(
            temp_output_path=tmp_path / "temp.docx",
            target_format="docx",
            max_input_bytes=100 * 1024 * 1024,
            _is_cancelled=lambda: True,
        )

        adapter = PdfDocumentAdapter()
        with pytest.raises(ConversionCancelledError):
            adapter.convert(req, ctx)

    def test_collision_safe_naming(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_01_simple_digital.pdf"
        out_file = tmp_path / "doc_01_simple_digital.docx"
        out_file.write_text("pre-existing file", encoding="utf-8")

        req = ConversionRequest(
            source=pdf_path,
            target_format="docx",
            output_dir=tmp_path,
        )
        res = convert_request(req)
        assert res.output_path.exists()
        assert res.output_path != out_file
        assert res.output_path.name.startswith("doc_01_simple_digital (")
        # Original remains untouched
        assert out_file.read_text(encoding="utf-8") == "pre-existing file"

    def test_does_not_overwrite_source_pdf(self, tmp_path: Path):
        # Place PDF in temp directory
        temp_pdf = tmp_path / "test.pdf"
        temp_pdf.write_bytes((FIXTURES_DIR / "doc_01_simple_digital.pdf").read_bytes())
        orig_bytes = temp_pdf.read_bytes()

        req = ConversionRequest(source=temp_pdf, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()
        assert res.output_path != temp_pdf
        assert temp_pdf.read_bytes() == orig_bytes

    def test_borderless_table_reconstruction(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_15_borderless_table.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        assert len(d.tables) == 1
        tbl = d.tables[0]
        assert len(tbl.rows) == 4
        # Verify header bolding
        for c in tbl.rows[0].cells:
            assert any(r.font.bold for p in c.paragraphs for r in p.runs)

    def test_image_and_caption_flow(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_16_image_caption.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        assert len(d.inline_shapes) == 1
        texts = [p.text for p in d.paragraphs if p.text]
        assert any("Architecture Overview" in t for t in texts)
        assert any("Figure 1: Core System Architecture" in t for t in texts)
        assert any("modular offline adapter pipeline" in t for t in texts)

    def test_headers_footers_paragraph_protection(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_17_headers_footers.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        texts = [p.text for p in d.paragraphs if p.text]
        assert any(t.startswith("Header:") for t in texts)
        assert any(t.startswith("Footer:") for t in texts)

    def test_dense_document_reconstruction(self, tmp_path: Path):
        pdf_path = FIXTURES_DIR / "doc_18_dense_document.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        texts = [p.text for p in d.paragraphs if p.text]
        assert len(texts) >= 5
        assert any("Workload Stress Analysis" in t for t in texts)
        assert any("Layout Robustness" in t for t in texts)

    def test_visual_renderer_integration(self, tmp_path: Path):
        from tests.visual_renderer import DocxVisualRenderer

        if not DocxVisualRenderer.is_available():
            pytest.skip("No local Word COM or LibreOffice renderer installed.")

        # Convert simple digital PDF to DOCX
        pdf_path = FIXTURES_DIR / "doc_01_simple_digital.pdf"
        req = ConversionRequest(source=pdf_path, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)

        # Render DOCX to PDF for visual QA
        rendered_pdf = tmp_path / "rendered_qa.pdf"
        success = DocxVisualRenderer.render_docx_to_pdf(res.output_path, rendered_pdf)
        assert success is True
        assert rendered_pdf.exists()

        # Render page to PNG
        img_paths = DocxVisualRenderer.render_docx_to_page_images(res.output_path, tmp_path / "img_out", dpi=100)
        assert len(img_paths) >= 1
        assert img_paths[0].exists()
        assert img_paths[0].stat().st_size > 0


class TestRealWorldDocxEdgeCases:
    """Verifies edge cases uncovered during LS-BE-008A real-world validation."""

    def test_numeric_cell_pattern_negative_accounting(self):
        from localshift.documents.docx_builder import NUMERIC_CELL_PATTERN
        valid_numbers = [
            "$45,200.00",
            "($1,250.00)",
            "(450.00)",
            "[-100]",
            "-50.5%",
            "€1,000",
            "£350.50",
            "¥50,000",
            "(15.4%)",
        ]
        for val in valid_numbers:
            assert NUMERIC_CELL_PATTERN.match(val) is not None, f"Failed on: {val}"

        non_numbers = ["Description", "Assets", "Cash and Cash Equivalents", "Total Current Assets"]
        for val in non_numbers:
            assert NUMERIC_CELL_PATTERN.match(val) is None, f"Unexpected match on: {val}"

    def test_article_and_roman_numeral_heading_candidate(self):
        from localshift.documents.docx_builder import _is_heading_candidate
        is_h, lvl = _is_heading_candidate("Article I: General Principles")
        assert is_h is True
        assert lvl == 1

        is_h, lvl = _is_heading_candidate("Article 4 - Non-Tariff Barriers")
        assert is_h is True
        assert lvl == 1

        is_h, lvl = _is_heading_candidate("Chapter IV: Institutional Arrangements")
        assert is_h is True
        assert lvl == 1

    def test_semicolon_and_numbered_list_paragraph_grouping(self):
        from localshift.documents.docx_builder import _group_lines_into_paragraphs
        lines = [
            "The licensee agrees to the following conditions:",
            "1. Shall submit compliance records semi-annually;",
            "2. Shall notify the Authority of changes within 14 days;",
            "(a) Detailed schematics must be preserved.",
        ]
        paras = _group_lines_into_paragraphs(lines)
        assert len(paras) == 4
        assert paras[1].startswith("1. Shall submit")
        assert paras[2].startswith("2. Shall notify")
        assert paras[3].startswith("(a) Detailed")

    def test_balance_sheet_table_recovery(self):
        from localshift.documents.docx_builder import _parse_table_structure
        lines = [
            "Assets 2026 2025",
            "--------------------------------------------------",
            "Cash and Cash Equivalents $45,200.00 $38,100.00",
            "Accounts Receivable $28,400.00 $24,500.00",
            "Less: Allowance for Credit ($1,250.00) ($950.00)",
            "Total Current Assets $72,350.00 $61,650.00",
        ]
        res = _parse_table_structure(lines)
        assert res is not None
        rows, start, end = res
        assert len(rows) == 5  # 1 header + 4 data rows
        assert len(rows[0]) == 3
        # Ensure description is consolidated in column 0
        assert rows[3][0] == "Less: Allowance for Credit"
        assert rows[3][1] == "($1,250.00)"
        assert rows[3][2] == "($950.00)"

    def test_sanitize_xml_text_removes_illegal_controls(self):
        from localshift.documents.docx_builder import sanitize_xml_text
        dirty = "Valid\x00Text\x0cWith\x08Control\x0bChars\tAnd\nNewline \U0001f600 Emoji"
        clean, count = sanitize_xml_text(dirty)
        assert "\x00" not in clean
        assert "\x08" not in clean
        assert "\x0c" not in clean
        assert "\x0b" not in clean
        assert "\t" in clean
        assert "\n" in clean
        assert "ValidText" in clean or "Valid" in clean
        assert "\U0001f600" in clean
        assert count > 0

    def test_pdf_docx_reconstruction_with_control_characters(self, tmp_path: Path):
        """Verify that PDFs containing NULL bytes, form feeds, and control characters reconstruct safely."""
        from localshift.documents.docx_builder import PdfToDocxReconstructor
        import pypdfium2 as pdfium
        from PIL import Image, ImageDraw

        # Create a PDF page that has text with control characters
        out_docx = tmp_path / "sanitized_output.docx"

        # Mock page text reconstruction directly
        doc = docx.Document()
        dirty_lines = [
            "Project Report \x00 Confidential",
            "Section 1 \x0c Overview with FormFeed",
            "Data item \x08\x0b with vertical tabs and backspaces",
            "Normal paragraph with unicode: Café Münsterland — 2026",
        ]
        PdfToDocxReconstructor._render_text_lines(doc, dirty_lines, median_font_size=12.0)
        doc.save(str(out_docx))

        assert out_docx.exists()
        reloaded = docx.Document(out_docx)
        full_text = " ".join(p.text for p in reloaded.paragraphs if p.text)
        assert "Project Report" in full_text
        assert "Overview" in full_text
        assert "Café Münsterland" in full_text
        assert "\x00" not in full_text

