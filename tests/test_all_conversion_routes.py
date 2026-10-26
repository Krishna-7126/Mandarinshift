"""
tests/test_all_conversion_routes.py — Exhaustive Verification Suite for all 54 Supported Conversion Routes.

Ensures:
1. Every single advertised conversion route in get_supported_routes() actually executes with real files.
2. Every output file is reopened and verified for structural integrity.
3. Source files remain strictly immutable (byte-for-byte identical).
4. Destination collision prevention preserves existing files.
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any
import pytest
from PIL import Image
import docx
import pypdfium2

from localshift import (
    ConversionRequest,
    convert_request,
    get_supported_routes,
)

ALL_ROUTES = get_supported_routes()


@pytest.fixture(scope="module")
def sample_sources(tmp_path_factory) -> dict[str, Path]:
    """Generate canonical valid test inputs for all supported source types."""
    dir_path = tmp_path_factory.mktemp("test_sources")
    sources: dict[str, Path] = {}

    # 1. Raster Images (RGB and RGBA)
    img_formats = ["png", "jpeg", "webp", "bmp", "tiff", "gif", "ico"]
    for fmt in img_formats:
        p = dir_path / f"source.{fmt}"
        mode = "RGBA" if fmt in ("png", "webp", "ico") else "RGB"
        color = (60, 120, 180, 200) if mode == "RGBA" else (60, 120, 180)
        im = Image.new(mode, (48, 48), color=color)
        im.save(p, format="PNG" if fmt == "png" else fmt.upper())
        sources[fmt] = p

    # 2. Rich Digital PDF
    fixtures_pdf = Path(__file__).parent / "fixtures" / "documents" / "doc_01_simple_digital.pdf"
    if fixtures_pdf.is_file():
        sources["pdf"] = fixtures_pdf
    else:
        pdf_path = dir_path / "source.pdf"
        doc = pypdfium2.PdfDocument.new()
        p = doc.new_page(300, 400)
        doc.save(pdf_path)
        doc.close()
        sources["pdf"] = pdf_path

    # 3. Valid Word DOCX
    docx_path = dir_path / "source.docx"
    d = docx.Document()
    d.add_heading("LocalShift Conversion Test", level=1)
    d.add_paragraph("Testing bidirectional document route.")
    d.save(docx_path)
    sources["docx"] = docx_path

    return sources


@pytest.mark.parametrize("source_format,target_format", ALL_ROUTES)
def test_conversion_route_execution_and_integrity(
    source_format: str,
    target_format: str,
    sample_sources: dict[str, Path],
    tmp_path: Path,
):
    """Verify that every advertised conversion route succeeds and creates structurally valid outputs."""
    src = sample_sources[source_format]
    src_bytes_before = src.read_bytes()

    out_dir = tmp_path / f"{source_format}_to_{target_format}"
    out_dir.mkdir(parents=True, exist_ok=True)

    req = ConversionRequest(
        source=src,
        target_format=target_format,
        output_dir=out_dir,
    )
    result = convert_request(req)

    # 1. Basic Result Checks
    assert result.output_path.exists(), f"Output file missing for route {source_format} -> {target_format}"
    assert result.output_path.stat().st_size > 0, f"Empty output produced for route {source_format} -> {target_format}"
    assert result.source_format == source_format
    assert result.target_format == target_format

    # 2. Source File Immutability Check
    assert src.read_bytes() == src_bytes_before, f"Source file was mutated during {source_format} -> {target_format}"

    # 3. Structural Integrity Verification (Re-opening target file)
    out_file = result.output_path
    if target_format in ("png", "jpeg", "webp", "bmp", "tiff", "gif", "ico"):
        with Image.open(out_file) as im:
            im.verify()
            assert im.size[0] > 0 and im.size[1] > 0

    elif target_format == "pdf":
        doc = pypdfium2.PdfDocument(out_file)
        assert len(doc) > 0
        doc.close()

    elif target_format == "docx":
        doc = docx.Document(out_file)
        assert len(doc.paragraphs) > 0 or len(doc.tables) > 0

    elif target_format == "txt":
        text = out_file.read_text(encoding="utf-8")
        assert len(text) > 0


def test_collision_avoidance_does_not_overwrite(sample_sources: dict[str, Path], tmp_path: Path):
    """Verify that existing output files are never overwritten and receive collision numbering."""
    src = sample_sources["png"]
    out_dir = tmp_path / "collision_test"
    out_dir.mkdir(parents=True, exist_ok=True)

    # First conversion
    req1 = ConversionRequest(source=src, target_format="jpeg", output_dir=out_dir)
    res1 = convert_request(req1)
    assert res1.output_path.name == "source.jpeg"

    # Second conversion with identical name & destination
    req2 = ConversionRequest(source=src, target_format="jpeg", output_dir=out_dir)
    res2 = convert_request(req2)
    assert res2.output_path.name == "source (2).jpeg"
    assert res1.output_path.exists()
    assert res2.output_path.exists()
