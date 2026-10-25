"""
tests/test_unseen_validation.py — Automated regression tests for unseen document validation set (LS-BE-007C).
"""
from __future__ import annotations

from pathlib import Path
import pytest

import docx
from localshift import ConversionRequest, convert_request
from tests.validation_generator import (
    UNSEEN_GROUND_TRUTH,
    generate_unseen_validation_fixtures,
)
from tests.visual_renderer import DocxVisualRenderer

VAL_FIXTURES_DIR = Path(__file__).parent / "fixtures" / "unseen_documents"


@pytest.fixture(scope="module", autouse=True)
def ensure_unseen_fixtures():
    """Ensure all 10 unseen validation fixtures exist before running tests."""
    return generate_unseen_validation_fixtures(VAL_FIXTURES_DIR)


class TestUnseenDocumentValidation:
    """Validate generalization of PDF -> DOCX reconstruction on unseen documents."""

    def test_unseen_fixtures_inventory(self):
        fixtures = list(VAL_FIXTURES_DIR.glob("*.pdf"))
        assert len(fixtures) == 10
        for f in fixtures:
            assert f.stat().st_size > 0

    def test_unseen_born_digital_contract_docx(self, tmp_path: Path):
        pdf = VAL_FIXTURES_DIR / "val_01_born_digital_contract.pdf"
        req = ConversionRequest(source=pdf, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        hdgs = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        assert len(hdgs) >= 2
        assert any("Grant of License" in h for h in hdgs)
        assert any("Limitation of Liability" in h for h in hdgs)

    def test_unseen_financial_ledger_table_recovery(self, tmp_path: Path):
        pdf = VAL_FIXTURES_DIR / "val_05_complex_financial_table.pdf"
        req = ConversionRequest(source=pdf, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        assert len(d.tables) == 1
        tbl = d.tables[0]
        assert len(tbl.rows) == 5
        assert len(tbl.columns) == 5

        # Verify header bolding
        for c in tbl.rows[0].cells:
            assert any(r.font.bold for p in c.paragraphs for r in p.runs)

        # Verify first row data cells
        row1_cells = [c.text for c in tbl.rows[1].cells]
        assert row1_cells[0] == "1010-01"
        assert row1_cells[1] == "Operating Cash Account"
        assert row1_cells[2] == "$45,200.00"

    def test_unseen_technical_manual_image_flow(self, tmp_path: Path):
        pdf = VAL_FIXTURES_DIR / "val_09_technical_manual_figures.pdf"
        req = ConversionRequest(source=pdf, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        assert len(d.inline_shapes) == 1
        texts = [p.text for p in d.paragraphs if p.text]
        assert any("System Architecture Manual" in t for t in texts)
        assert any("Figure 1: Core Processing Pipeline" in t for t in texts)

    def test_unseen_multilingual_legal_unicode(self, tmp_path: Path):
        pdf = VAL_FIXTURES_DIR / "val_07_multilingual_legal.pdf"
        req = ConversionRequest(source=pdf, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        full_text = "\n".join(p.text for p in d.paragraphs)
        glyphs = ["schützt", "persönliche", "Déclaration", "conformité", "Garantía", "Bekreftelse", "€", "£", "¥", "$"]
        for g in glyphs:
            assert g in full_text

    def test_unseen_long_dense_manuscript_pagination(self, tmp_path: Path):
        pdf = VAL_FIXTURES_DIR / "val_10_long_dense_manuscript.pdf"
        req = ConversionRequest(source=pdf, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)
        assert res.output_path.exists()

        d = docx.Document(res.output_path)
        hdgs = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        assert len(hdgs) == 5

    def test_unseen_visual_qa_rendering(self, tmp_path: Path):
        if not DocxVisualRenderer.is_available():
            pytest.skip("No local Word COM or LibreOffice renderer installed.")

        pdf = VAL_FIXTURES_DIR / "val_01_born_digital_contract.pdf"
        req = ConversionRequest(source=pdf, target_format="docx", output_dir=tmp_path)
        res = convert_request(req)

        rendered_pdf = tmp_path / "contract_qa.pdf"
        success = DocxVisualRenderer.render_docx_to_pdf(res.output_path, rendered_pdf)
        assert success is True
        assert rendered_pdf.exists()
        assert rendered_pdf.stat().st_size > 0
